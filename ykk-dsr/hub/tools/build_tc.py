#!/usr/bin/env python3
"""Build a TC workbook (+ TestData workbook) from a JSON spec, using the hub templates.

Usage: python3 -I ykk-dsr/hub/tools/build_tc.py spec.json OUT_DIR

spec.json:
{
  "us": "DSR-566", "folder_id": "357", "us_name": "[レイアウト] ニュース詳細（news-article）",
  "jira_url": "https://sutrix.atlassian.net/browse/DSR-357", "module": "ニュース詳細 / News detail",
  "tcs": [ {"id": "1.0", "feature": "...", "summary": "EN: ... JP: ...", "pre": "EN: ... JP: ...",
            "steps": "EN: 1. ... JP: 1. ...", "expected": "EN: ... JP: ...", "design": "BAC#1",
            "layer": "Page", "exec": "Automation", "lang": "ja, en, cn", "device": "PC Chromium",
            "td": "TD-566-01", "blocked": "—"} ],
  "assumptions": [ ["A1", "text", "source", "impact"] ],
  "testdata": [ ["TD-566-01", "1.0", "page", "/master/ja/...", "field", "ja", "en", "zh", "Yes"] ],
  "howto": [ ["1", "step", "check / if X then Y"] ]
}
Output files: "[Sutrix] DSR -<folder_id>- <us_name>.xlsx" and "..._TestData.xlsx".
Exec type must be one of: Automation | Visual (Playwright VPS) | Manual.
"""
import json, os, sys, shutil
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, "..", "templates")
EXEC = {"Automation", "Visual (Playwright VPS)", "Manual"}


def build(spec, out_dir):
    name = f"[Sutrix] DSR -{spec['folder_id']}- {spec['us_name']}"
    tc_path = os.path.join(out_dir, name + ".xlsx")
    td_path = os.path.join(out_dir, name + "_TestData.xlsx")
    wb = openpyxl.load_workbook(os.path.join(TPL, "TC_TEMPLATE.xlsx"))
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=1, max_row=6):
        for c in row:
            if isinstance(c.value, str) and "browse/DSR-<ID>" in c.value:
                c.value = spec["jira_url"]
    ws["A11"] = spec.get("module", spec["us_name"])
    style_row = [ws.cell(12, c) for c in range(1, 22)]
    for i, tc in enumerate(spec["tcs"]):
        if tc.get("exec") not in EXEC:
            raise SystemExit(f"bad exec type in {tc.get('id')}: {tc.get('exec')}")
        r = 12 + i
        vals = {1: tc["id"], 2: tc["feature"], 3: tc["summary"], 4: tc["pre"], 5: tc["steps"],
                6: tc["expected"], 8: "To do / 未実施", 9: "To do / 未実施", 10: "To do / 未実施",
                12: tc.get("design", ""), 16: tc["layer"], 17: tc["exec"], 18: tc["lang"],
                19: tc["device"], 20: tc["td"], 21: tc.get("blocked", "—")}
        for c in range(1, 22):
            cell = ws.cell(r, c)
            if r != 12:
                src = style_row[c - 1]
                cell.font, cell.border, cell.fill = src.font.copy(), src.border.copy(), src.fill.copy()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.value = vals.get(c)
    wa = wb["Assumptions"]
    for i, a in enumerate(spec.get("assumptions", []), start=2):
        for j, v in enumerate(a, start=1):
            wa.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(tc_path)

    t = openpyxl.load_workbook(os.path.join(TPL, "TESTDATA_TEMPLATE.xlsx"))
    d = t["Data"]
    for i, row in enumerate(spec.get("testdata", []), start=2):
        for j, v in enumerate(row, start=1):
            d.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    h = t["How-to-create"]
    for r in range(2, h.max_row + 1):
        for c in range(1, 4):
            h.cell(r, c).value = None
    for i, row in enumerate(spec.get("howto", []), start=2):
        for j, v in enumerate(row, start=1):
            h.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    t.save(td_path)
    print(tc_path); print(td_path)


if __name__ == "__main__":
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    os.makedirs(sys.argv[2], exist_ok=True)
    build(spec, sys.argv[2])

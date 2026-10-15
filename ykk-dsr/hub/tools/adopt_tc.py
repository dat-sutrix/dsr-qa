#!/usr/bin/env python3
"""Adopt an existing dev TC workbook: keep dev content, add QA columns P–U, Assumptions sheet,
optional QA-added TC rows, label fixes; build TestData from the hub template.

Usage: python3 -I ykk-dsr/hub/tools/adopt_tc.py spec.json DEV_TC.xlsx OUT_TC.xlsx templates/TESTDATA_TEMPLATE.xlsx OUT_TestData.xlsx
spec keys: label_fixes {cell: "new value" | [old, new]}, qa_cols {tc_id: [layer, exec, lang, device, td, blocked]},
           added_tcs [{id, feature, summary, pre, steps, expected, design}], assumptions, testdata, howto.
Multi-sheet: "sheets": [{sheet: <index>, label_fixes, qa_cols, added_tcs}, ...] instead of the 3 top-level keys.
Example spec: ykk-dsr/hub/tc-out/DSR-568.spec.json (single), DSR-608.spec.json (multi-sheet)
"""
import json, sys, copy
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
spec = json.load(open(sys.argv[1], encoding="utf-8")); src, dst = sys.argv[2], sys.argv[3]
wb = openpyxl.load_workbook(src)
al = Alignment(wrap_text=True, vertical="top")


def adopt_sheet(ws, sp):
    hdr = ws["O10"]
    for i, h in enumerate(["Layer (CF/Page)", "Exec type", "Language", "Device / Browser", "Test data (TD ID)", "Blocked-by"]):
        c = ws.cell(10, 16 + i, h); c.font = copy.copy(hdr.font); c.fill = copy.copy(hdr.fill); c.border = copy.copy(hdr.border); c.alignment = copy.copy(hdr.alignment)
        ws.column_dimensions[c.column_letter].width = 22
    for cell, fx in sp["label_fixes"].items():
        ws[cell] = ws[cell].value.replace(*fx) if isinstance(fx, list) else fx
    rows = {str(ws.cell(r, 1).value): r for r in range(12, ws.max_row + 1) if ws.cell(r, 1).value}
    last = max(rows.values())
    for tc in sp["added_tcs"]:
        last += 1; rows[tc["id"]] = last
        for c in range(1, 16):
            s = ws.cell(last - 1, c); d = ws.cell(last, c)
            d.font, d.border, d.fill = copy.copy(s.font), copy.copy(s.border), copy.copy(s.fill); d.alignment = al
        vals = {1: float(tc["id"]), 2: tc["feature"], 3: tc["summary"], 4: tc["pre"], 5: tc["steps"], 6: tc["expected"],
                8: "To do / 未実施", 9: "To do / 未実施", 10: "To do / 未実施", 12: tc["design"]}
        for k, v in vals.items(): ws.cell(last, k, v)
        ws.row_dimensions[last].height = ws.row_dimensions[last - 1].height
    for dv in ws.data_validations.dataValidation:
        if str(dv.sqref).startswith("H12:J"): dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"H12:J{last}")
    ref = ws.cell(12, 15)
    for tid, vals in sp["qa_cols"].items():
        r = rows[tid]
        for i, v in enumerate(vals):
            c = ws.cell(r, 16 + i, v); c.alignment = al; c.border = copy.copy(ref.border); c.font = copy.copy(ref.font)
    assert set(sp["qa_cols"]) == set(rows), (set(rows) ^ set(sp["qa_cols"]))
    return len(rows)


sheets = spec.get("sheets") or [dict(sheet=0, label_fixes=spec["label_fixes"], qa_cols=spec["qa_cols"], added_tcs=spec["added_tcs"])]
nrows = sum(adopt_sheet(wb.worksheets[sp["sheet"]], sp) for sp in sheets)
wa = wb.create_sheet("Assumptions")
for j, h in enumerate(["ID", "Assumption / 前提", "Source", "Impact if wrong"], 1):
    c = wa.cell(1, j, h); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E1F2")
for i, a in enumerate(spec["assumptions"], 2):
    for j, v in enumerate(a, 1): wa.cell(i, j, v).alignment = al
for col, w in zip("ABCD", (6, 90, 35, 30)): wa.column_dimensions[col].width = w
wb.save(dst); print(dst, "sheets", len(sheets), "rows", nrows)

t = openpyxl.load_workbook(sys.argv[4]); d = t["Data"]
for i, row in enumerate(spec["testdata"], 2):
    for j, v in enumerate(row, 1): d.cell(i, j, v).alignment = al
h = t["How-to-create"]
for r in range(2, h.max_row + 1):
    for c in range(1, 4): h.cell(r, c).value = None
for i, row in enumerate(spec["howto"], 2):
    for j, v in enumerate(row, 1): h.cell(i, j, v).alignment = al
t.save(sys.argv[5]); print(sys.argv[5])
